"""Post-evaluation consistency audit; does not replace official verifier scores."""
import json
import statistics
from pathlib import Path

import openpyxl

root = Path('/audit')
path = root / 'benchmark/tasks/xlsx-recover-data/environment/nasa_budget_incomplete.xlsx'
book = openpyxl.load_workbook(path, data_only=True)
budget = book['Budget by Directorate']
growth = book['Growth Analysis']
science = []
for row in range(8, 14):
    value = budget.cell(row, 2).value
    if value == '???':
        value = budget.cell(row, 11).value - sum(budget.cell(row, col).value for col in range(3, 11))
    science.append(value)
checks = []
for col in range(3, 11):
    values = [budget.cell(row, col).value for row in range(8, 14)]
    known = growth.cell(8, col).value
    if all(isinstance(value, (int, float)) for value in values) and isinstance(known, (int, float)):
        checks.append({'column': openpyxl.utils.get_column_letter(col), 'supplied_average': known,
                       'five_year_average': round(statistics.mean(values[:5]), 1),
                       'six_year_average': round(statistics.mean(values), 1)})
data = {
    'purpose': 'Independent audit after all evaluation runs; original scores remain unchanged',
    'growth_labels': [[growth.cell(row, col).value for col in range(1, 4)] for row in [1, 3, 4, 5, 6, 7, 8]],
    'science_2019_to_2024': science,
    'science_five_year_average': round(statistics.mean(science[:5]), 1),
    'science_six_year_average': round(statistics.mean(science), 1),
    'other_supplied_averages': checks,
    'official_expected_science_average': 7444.4,
    'all_four_agents_science_average': 7610.3,
    'finding': 'The oracle averages five science entries (2019–2023), while the other supplied average-budget entries follow the six-entry 2019–2024 window. All agents use that six-entry convention. This is a reference/data consistency issue, not evidence that learned skills improved or degraded accuracy.',
}
(root / 'benchmark-audit.json').write_text(json.dumps(data, indent=2) + '\n')
print(json.dumps(data, ensure_ascii=False))
