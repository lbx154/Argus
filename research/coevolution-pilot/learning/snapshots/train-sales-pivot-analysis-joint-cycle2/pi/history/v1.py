#!/usr/bin/env python3
"""Bounded, read-only inspection of CSV, XLSX, and text PDFs."""
import os
# Keep optional numerical backends imported by spreadsheet libraries from
# creating large thread pools in constrained inspection environments.
for key in ('OPENBLAS_NUM_THREADS', 'OMP_NUM_THREADS', 'MKL_NUM_THREADS', 'NUMEXPR_NUM_THREADS'):
    os.environ[key] = '1'
import argparse
import csv
import json
import sys
from pathlib import Path

MAX_TEXT = 4000


def clipped(value):
    if isinstance(value, str) and len(value) > MAX_TEXT:
        return value[:MAX_TEXT] + '…'
    return value


def inspect_csv(path, limit):
    with path.open('r', encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        columns = reader.fieldnames or []
        items = []
        for row_no, row in enumerate(reader, 2):
            if len(items) >= limit:
                break
            items.append({'source_id': f'row:{row_no}', 'row': {name: clipped(row.get(name, '')) for name in columns}})
    return {'kind': 'csv', 'items': items, 'columns': columns}


def inspect_xlsx(path, limit):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    items = []
    try:
        sheets = list(workbook.sheetnames)
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            columns = [str(value) if value is not None else f'column_{index}' for index, value in enumerate(header, 1)]
            for row_no, values in enumerate(rows, 2):
                if len(items) >= limit:
                    break
                record = {columns[index]: clipped(values[index] if index < len(values) else None) for index in range(len(columns))}
                items.append({'source_id': f'{sheet.title}!{row_no}', 'sheet': sheet.title, 'row': record})
            if len(items) >= limit:
                break
    finally:
        workbook.close()
    return {'kind': 'xlsx', 'items': items, 'sheets': sheets}


def inspect_pdf(path, limit):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    items = []
    for page_no, page in enumerate(reader.pages, 1):
        if len(items) >= limit:
            break
        items.append({'source_id': f'page:{page_no}', 'page': page_no, 'text': clipped(page.extract_text() or '')})
    return {'kind': 'pdf', 'items': items, 'pages': len(reader.pages)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path')
    parser.add_argument('--limit', type=int, default=20)
    args = parser.parse_args()
    path = Path(args.path)
    if args.limit < 0:
        parser.error('--limit must be non-negative')
    if not path.is_file():
        print(json.dumps({'kind': 'error', 'items': [], 'path': str(path), 'error': 'file not found'}, ensure_ascii=False))
        return 2
    readers = {'.csv': inspect_csv, '.xlsx': inspect_xlsx, '.pdf': inspect_pdf}
    reader = readers.get(path.suffix.lower())
    if reader is None:
        print(json.dumps({'kind': 'error', 'items': [], 'path': str(path), 'error': 'unsupported format'}, ensure_ascii=False))
        return 2
    try:
        result = reader(path, args.limit)
        result['path'] = str(path)
        print(json.dumps(result, ensure_ascii=False, default=str))
        return 0
    except Exception as exc:
        print(json.dumps({'kind': 'error', 'items': [], 'path': str(path), 'error': f'{type(exc).__name__}: {exc}'}, ensure_ascii=False))
        return 1


if __name__ == '__main__':
    sys.exit(main())
